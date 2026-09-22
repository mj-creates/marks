import os
import re
import shutil
from typing import Dict, List, Tuple
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def clean_component_name(raw_comp: str) -> str:
    """Format assessment components into clean, standardized title like 'Week Test-1 Max:10'."""
    comp = raw_comp.strip()
    comp = re.sub(r'Max\s*(\d+)', r'Max:\1', comp, flags=re.I)
    comp = re.sub(r'\bweek test', 'Week Test', comp, flags=re.I)
    comp = re.sub(r'\bmid', 'Mid', comp, flags=re.I)
    comp = re.sub(r'\bexperiment', 'Experiment', comp, flags=re.I)
    comp = re.sub(r'\binternal lab', 'Internal Lab', comp, flags=re.I)
    comp = re.sub(r'\bminor project', 'Minor Project', comp, flags=re.I)
    comp = re.sub(r'(\d)(Max:)', r'\1 \2', comp)
    return comp

def natural_sort_key(s: str) -> list:
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', str(s))]

def style_merged_range(ws, min_r, min_c, max_r, max_c, font=None, fill=None, align=None, border=None):
    for r in range(min_r, max_r + 1):
        for c in range(min_c, max_c + 1):
            cell = ws.cell(row=r, column=c)
            if font: cell.font = font
            if fill: cell.fill = fill
            if align: cell.alignment = align
            if border: cell.border = border

def build_hierarchical_master_sheet(ws, df_master: pd.DataFrame, subject_map: Dict[str, List[Tuple[str, str]]]):
    """
    Builds a 2-tier Hierarchical Master Excel Sheet matching Image 1:
      Row 1: [REGD.NO (merged A1:A2), NAME (merged B1:B2), SECTION (merged C1:C2),
              BCI (merged over BCI cols), BEEE (merged over BEEE cols)...]
      Row 2: [Week Test-1 Max:10, Week Test-2 Max:10, Mid-1 Max:30...]
      Row 3+: Student Data rows
    """
    # Palette definitions
    top_hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # Dark Navy Blue
    sub_hdr_fill = PatternFill(start_color="285E8E", end_color="285E8E", fill_type="solid")  # Medium Navy Blue
    top_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    sub_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    
    zebra_fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="Calibri", size=10, bold=False, color="000000")

    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align   = Alignment(horizontal="left", vertical="center")
    right_align  = Alignment(horizontal="right", vertical="center")

    thin_side = Side(style="thin", color="D9D9D9")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 28

    # 1. Primary Identifier Headers (merged rows 1:2)
    col_id_labels = [(1, "REGD.NO"), (2, "NAME"), (3, "SECTION")]
    for col_idx, label in col_id_labels:
        ws.merge_cells(start_row=1, start_column=col_idx, end_row=2, end_column=col_idx)
        ws.cell(row=1, column=col_idx, value=label)
        style_merged_range(ws, 1, col_idx, 2, col_idx, font=top_font, fill=top_hdr_fill, align=center_align, border=cell_border)

    # 2. Subject Groups and Assessment Headers
    current_col = 4
    for subj in sorted(subject_map.keys(), key=natural_sort_key):
        items = subject_map[subj]
        start_col = current_col
        end_col = current_col + len(items) - 1

        # Merge Row 1 across this subject's columns
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        ws.cell(row=1, column=start_col, value=subj.upper())
        style_merged_range(ws, 1, start_col, 1, end_col, font=top_font, fill=top_hdr_fill, align=center_align, border=cell_border)

        # Set Row 2 for each assessment component under this subject
        for raw_u, comp_clean in items:
            cell_r2 = ws.cell(row=2, column=current_col, value=comp_clean)
            cell_r2.font = sub_font
            cell_r2.fill = sub_hdr_fill
            cell_r2.alignment = center_align
            cell_r2.border = cell_border
            current_col += 1

    # 3. Data Rows starting at Row 3
    num_rows = len(df_master)
    total_cols = current_col - 1

    for row_idx, row_data in enumerate(df_master.itertuples(index=False), start=3):
        ws.row_dimensions[row_idx].height = 20
        row_fill = white_fill if (row_idx % 2 == 1) else zebra_fill

        row_vals = ["" if pd.isna(v) else v for v in row_data]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill
            cell.font = data_font
            cell.border = cell_border

            if col_idx == 1:
                cell.alignment = center_align
            elif col_idx == 2:
                cell.alignment = left_align
            elif col_idx == 3:
                cell.alignment = center_align
            else:
                if isinstance(val, (int, float)):
                    cell.alignment = right_align
                else:
                    cell.alignment = center_align

    # Freeze Panes at D3 (Row 1, Row 2, and Columns A, B, C remain locked!)
    ws.freeze_panes = "D3"

    # Column Widths
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 12

    col_idx = 4
    for subj in sorted(subject_map.keys(), key=natural_sort_key):
        for raw_u, comp_clean in subject_map[subj]:
            col_letter = get_column_letter(col_idx)
            hdr_len = len(comp_clean)
            sample_lens = [len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(3, min(num_rows + 3, 25))]
            max_data = max(sample_lens) if sample_lens else 0
            ws.column_dimensions[col_letter].width = max(12, min(36, max(hdr_len, max_data) + 4))
            col_idx += 1

def style_single_sheet(ws, df: pd.DataFrame):
    """Standard single-header sheet styling."""
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    zebra_fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="Calibri", size=10, bold=False, color="000000")
    
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_side = Side(style="thin", color="D9D9D9")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    headers = list(df.columns)
    ws.append(headers)
    ws.row_dimensions[1].height = 32

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = cell_border

    for row_idx, row_data in enumerate(df.itertuples(index=False), start=2):
        ws.row_dimensions[row_idx].height = 20
        row_fill = white_fill if (row_idx % 2 == 0) else zebra_fill

        row_vals = ["" if pd.isna(v) else v for v in row_data]
        ws.append(row_vals)

        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = row_fill
            cell.font = data_font
            cell.border = cell_border

            if col_idx == 1:
                cell.alignment = center_align
            elif col_idx == 2:
                cell.alignment = left_align
            elif col_idx == 3:
                cell.alignment = center_align
            else:
                if isinstance(val, (int, float)):
                    cell.alignment = right_align
                else:
                    cell.alignment = center_align

    ws.freeze_panes = "D2"

    for col_idx, col_name in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        if col_name == "REGD.NO":
            ws.column_dimensions[col_letter].width = 16
        elif col_name == "NAME":
            ws.column_dimensions[col_letter].width = 30
        elif col_name == "SECTION":
            ws.column_dimensions[col_letter].width = 12
        else:
            hdr_len = len(str(col_name))
            sample_lens = [len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(2, min(len(df) + 2, 25))]
            max_data = max(sample_lens) if sample_lens else 0
            ws.column_dimensions[col_letter].width = max(12, min(36, max(hdr_len, max_data) + 4))

def run_consolidation():
    raw_path = r"c:\Users\spjee\marks\portal project\output\master_marks_batch2020_y1s1_20260921_162851.xlsx"
    out_dir = r"c:\Users\spjee\marks\portal project\output"
    out_file = os.path.join(out_dir, "Consolidated_Master_Marks.xlsx")
    workspace_root_file = r"c:\Users\spjee\marks\Consolidated_Master_Marks.xlsx"
    downloads_file = os.path.expanduser(r"~\Downloads\Consolidated_Master_Marks.xlsx")
    recent_user_file = os.path.join(out_dir, "master_marks_batch2020_y1s1_20260921_173452.xlsx")

    print(f"Loading raw dataset from:\n  {raw_path}")
    df_raw = pd.read_excel(raw_path)

    # 1. Primary Identifiers
    regd_col = "Regd No"
    sec_col = "Section"
    assert regd_col in df_raw.columns, f"Column '{regd_col}' not found"

    name_cols = [c for c in df_raw.columns if "NAME" in str(c).upper()]
    print(f"Detected {len(name_cols)} section NAME columns: {name_cols}")

    def extract_name(row):
        for c in name_cols:
            val = row[c]
            if pd.notna(val) and str(val).strip():
                return str(val).strip().upper()
        return ""

    names_series = df_raw.apply(extract_name, axis=1)

    # 2. Header Unification & Cleaning
    prefix_regex = re.compile(r'^VFSTR\s*B\.?\s*TECH.*?\d+[-_ ]Section\s*[-_ ]*', re.I)
    ignored_cols = set([regd_col, sec_col, "Overall - EXT", "overall - ext"] + name_cols)

    col_mapping = {}
    for col in df_raw.columns:
        if col in ignored_cols:
            continue
        clean_hdr = prefix_regex.sub('', str(col)).strip()
        if clean_hdr:
            col_mapping[col] = clean_hdr

    raw_unified_headers = sorted(list(set(col_mapping.values())), key=natural_sort_key)

    # Consolidate values for each unified header
    unified_values = {}
    for u_hdr in raw_unified_headers:
        matching_cols = [c for c, u in col_mapping.items() if u == u_hdr]
        sub_df = df_raw[matching_cols]
        unified_values[u_hdr] = sub_df.bfill(axis=1).iloc[:, 0]

    # Organize by Subject: mapping subject -> list of (raw_unified_hdr, clean_component_name)
    subject_map: Dict[str, List[Tuple[str, str]]] = {}
    for u_hdr in raw_unified_headers:
        parts = u_hdr.split(' ', 1)
        subj = parts[0]
        comp_raw = parts[1] if len(parts) > 1 else ""
        comp_clean = clean_component_name(comp_raw)
        if subj not in subject_map:
            subject_map[subj] = []
        subject_map[subj].append((u_hdr, comp_clean))

    print(f"Subjects detected ({len(subject_map)}): {list(subject_map.keys())}")

    # Build Master DataFrame
    master_dict = {
        "REGD.NO": df_raw[regd_col].astype(str).str.strip().str.upper(),
        "NAME": names_series,
        "SECTION": df_raw[sec_col].astype(str).str.strip()
    }
    
    for subj in sorted(subject_map.keys(), key=natural_sort_key):
        for raw_u, comp_clean in subject_map[subj]:
            col_label = f"{subj} - {comp_clean}"
            master_dict[col_label] = unified_values[raw_u]

    df_master = pd.DataFrame(master_dict)
    df_master = df_master.sort_values(by="REGD.NO", ascending=True).reset_index(drop=True)

    # Build Multi-Sheet Workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    # 1. Primary Sheet: 2-tier Hierarchical Master (Matches Image 1 Header + All Subjects)
    ws_master = wb.create_sheet(title="Master (All Subjects)")
    build_hierarchical_master_sheet(ws_master, df_master, subject_map)

    # 2. Individual Subject Tabs
    for subj in sorted(subject_map.keys(), key=natural_sort_key):
        subj_tab_name = subj[:31]
        ws_subj = wb.create_sheet(title=subj_tab_name)

        subj_dict = {
            "REGD.NO": df_master["REGD.NO"],
            "NAME": df_master["NAME"],
            "SECTION": df_master["SECTION"],
        }
        for raw_u, comp_clean in subject_map[subj]:
            subj_dict[comp_clean] = unified_values[raw_u]

        df_subj = pd.DataFrame(subj_dict)
        mark_cols = [c for c in df_subj.columns if c not in ["REGD.NO", "NAME", "SECTION"]]
        has_marks = df_subj[mark_cols].notna().any(axis=1)
        df_subj_filtered = df_subj[has_marks].sort_values(by="REGD.NO").reset_index(drop=True)
        style_single_sheet(ws_subj, df_subj_filtered)

    # 3. Flat Single-Row Master Tab
    ws_flat = wb.create_sheet(title="Master (Single Header)")
    style_single_sheet(ws_flat, df_master)

    # Save to deliverables
    wb.save(out_file)
    print(f"Saved consolidated master to: {out_file}")

    for dest in [workspace_root_file, downloads_file, recent_user_file]:
        try:
            shutil.copyfile(out_file, dest)
            print(f"Successfully updated: {dest}")
        except Exception as e:
            print(f"Could not update {dest}: {e}")

    print(f"\nCompleted! Workbook now contains {len(wb.sheetnames)} sheets:")
    for s in wb.sheetnames:
        print(f"  - {s}")

if __name__ == "__main__":
    run_consolidation()
